"""One-off script: insert live field-verification photos as Figures 7-10
into Appendix B, and register them in the List of Figures. Run once, then
rebuild the TOC via the existing Word COM step."""
import sys
sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

MANUSCRIPT = "../docs/Astorga_Central_Manuscript.docx"
SHOTS = "../docs/screenshots"

FIGURES = [
    ("live-setup.jpg", "Figure 7. Live Field Deployment Setup (ESP8266 Sensor Node, JSN-SR04T Probe, and Portable Power Supply)."),
    ("live-normal.jpg", "Figure 8. Live Field Verification — Normal Status (Barangay Hall Bridge node, 0 cm water level)."),
    ("live-watch.jpg", "Figure 9. Live Field Verification — Watch Status (Barangay Hall Bridge node, 44 cm water level, rising)."),
    ("live-critical.jpg", "Figure 10. Live Field Verification — Critical Status (Barangay Hall Bridge node, 77 cm water level, flooding)."),
]

BODY_TEXT = (
    "Figures 7 to 10 document a live end-to-end field verification of the "
    "prototype, using the physically deployed ESP8266 sensor node reporting "
    "over a mobile hotspot to the operations dashboard in real time. The "
    "water level at the Barangay Hall Bridge choke point was manually varied "
    "to walk the node through all three classification states — Normal, "
    "Watch, and Critical — with the dashboard reflecting each transition "
    "within the system's 1 Hz reporting interval."
)


def insert_paragraph_before(reference_paragraph, text=None, style=None):
    new_p = OxmlElement("w:p")
    reference_paragraph._p.addprevious(new_p)
    from docx.text.paragraph import Paragraph
    new_para = Paragraph(new_p, reference_paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def main():
    doc = Document(MANUSCRIPT)
    paragraphs = doc.paragraphs

    # --- 1. List of Figures: insert entries 7-10 right after Figure 6 entry ---
    lof_anchor = None
    for p in paragraphs:
        if p.text.strip() == "Figure 6. Instant Push Alert (Discord)":
            lof_anchor = p
            break
    if lof_anchor is None:
        raise RuntimeError("Could not find List of Figures anchor (Figure 6 entry)")

    lof_captions = [
        "Figure 7. Live Field Deployment Setup (ESP8266 Sensor Node)",
        "Figure 8. Live Field Verification — Normal Status",
        "Figure 9. Live Field Verification — Watch Status",
        "Figure 10. Live Field Verification — Critical Status",
    ]
    # Find the paragraph immediately after the anchor to insert before (so
    # order stays 6,7,8,9,10 rather than reversed).
    body = doc.element.body
    anchor_idx = list(body).index(lof_anchor._p)
    next_p_element = list(body)[anchor_idx + 1]
    from docx.text.paragraph import Paragraph
    next_para = Paragraph(next_p_element, lof_anchor._parent)
    for cap in lof_captions:
        insert_paragraph_before(next_para, cap, style=lof_anchor.style)

    # --- 2. Appendix B: insert new figures right before "Appendix C" ---
    appendix_c = None
    for p in doc.paragraphs:
        if p.text.strip() == "Appendix C: Cost of Materials":
            appendix_c = p
            break
    if appendix_c is None:
        raise RuntimeError("Could not find Appendix C heading")

    # blank spacer before the new block
    insert_paragraph_before(appendix_c, "", style="Normal")

    for filename, caption in FIGURES:
        img_para = insert_paragraph_before(appendix_c, style="Normal")
        img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = img_para.add_run()
        run.add_picture(f"{SHOTS}/{filename}", width=Inches(6.0))

        cap_para = insert_paragraph_before(appendix_c, style="Normal")
        cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap_para.add_run(caption)
        cap_run.italic = True
        cap_run.font.size = Pt(10)

    body_para = insert_paragraph_before(appendix_c, BODY_TEXT, style="Normal")
    insert_paragraph_before(appendix_c, "", style="Normal")

    doc.save(MANUSCRIPT)
    print("Saved:", MANUSCRIPT)


if __name__ == "__main__":
    main()
